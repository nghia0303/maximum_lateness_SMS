"""Validate every solver result against the authoritative filename lists.

The validator does not rerun any solver and does not modify source workbooks or
solution files. It checks persisted schedules against the original instances,
recomputes true Lmax, and writes a separate Excel report.
"""

from pathlib import Path
import math
import re
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from common.dataset import read_dataset
from common.project_paths import DATASET_SIZES, FILENAMES_DIR, INS_DIR, INSTANCE_TYPES, RESULTS_DIR
from common.schedule_utils import compute_max_lateness, validate_schedule


DATA_DIR = PROJECT_ROOT / "2016"
REPORT_WORKBOOK = RESULTS_DIR / "validation_report.xlsx"

SOLVER_SOURCES = {
    "DIRECTSAT": ("directSAT", RESULTS_DIR / "results_directSAT.xlsx"),
    "PROPOSED SAT": ("proposedSAT", RESULTS_DIR / "results_proposedSAT.xlsx"),
    "MP": ("cplex_mp", RESULTS_DIR / "results_cplex_mp.xlsx"),
    "GUROBI": ("gurobi", RESULTS_DIR / "results_gurobi.xlsx"),
    "CPLEX CP": ("cplex_cp", RESULTS_DIR / "results_cplex_cp.xlsx"),
    "CP_SAT": ("cpsat", RESULTS_DIR / "results_cpsat.xlsx"),
}

LATE_LINE_RE = re.compile(r"Lmax\s*=\s*(-?\d+)")
JOB_LINE_RE = re.compile(r"Job\s+(\d+):\s+start\s*=\s*(-?\d+)")
TERMINAL_STATUSES = {"UNSAT", "TIMEOUT", "INFEASIBLE", "TIME_LIMIT_FEASIBLE"}


def as_int(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value or value in {"-", "None", "nan"}:
            return None
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return None


def read_filename_instances():
    rows = []
    for size in DATASET_SIZES:
        for instance_type in INSTANCE_TYPES:
            type_specific = FILENAMES_DIR / f"{size}-{instance_type}.txt"
            filelist = type_specific if type_specific.exists() else FILENAMES_DIR / f"{size}.txt"
            if not filelist.exists():
                raise FileNotFoundError(f"Filename list not found: {filelist}")
            for line in filelist.read_text(encoding="utf-8").splitlines():
                filename = line.strip()
                if not filename:
                    continue
                rows.append(
                    {
                        "n": size,
                        "type": instance_type,
                        "filename": filename,
                    }
                )
    return rows


def read_solver_results(workbook_path):
    rows = {}
    if not workbook_path.exists():
        return rows

    for sheet_name in pd.ExcelFile(workbook_path).sheet_names:
        dataframe = pd.read_excel(workbook_path, sheet_name=sheet_name)
        for source_row in dataframe.to_dict("records"):
            size = as_int(source_row.get("n"))
            instance_type = source_row.get("type")
            if size is None or pd.isna(instance_type):
                dataset = str(source_row.get("dataset", sheet_name))
                size_text, instance_type = dataset.split("-", 1)
                size = int(size_text)
            filename = str(source_row.get("filename"))
            rows[(size, str(instance_type), filename)] = {
                "value": as_int(source_row.get("Lmax")),
                "status": str(source_row.get("status")),
            }
    return rows


def parse_solution_file(solution_path):
    if not solution_path.exists():
        return {
            "exists": False,
            "marker": "MISSING_FILE",
            "lmax": None,
            "schedule": {},
            "duplicate_jobs": 0,
        }

    lines = [line.strip() for line in solution_path.read_text(encoding="utf-8").splitlines() if line.strip()]
    if not lines:
        return {
            "exists": True,
            "marker": "EMPTY_FILE",
            "lmax": None,
            "schedule": {},
            "duplicate_jobs": 0,
        }

    first_line = lines[0]
    if first_line in TERMINAL_STATUSES:
        marker = first_line
    elif first_line.startswith("STATUS_"):
        marker = first_line
    elif LATE_LINE_RE.search(first_line):
        marker = "LMAX_FIRST"
    else:
        marker = first_line

    lmax = None
    schedule = {}
    duplicate_jobs = 0
    for line in lines:
        lmax_match = LATE_LINE_RE.search(line)
        if lmax_match and lmax is None:
            lmax = int(lmax_match.group(1))
        job_match = JOB_LINE_RE.search(line)
        if job_match:
            job = int(job_match.group(1))
            if job in schedule:
                duplicate_jobs += 1
            schedule[job] = int(job_match.group(2))

    return {
        "exists": True,
        "marker": marker,
        "lmax": lmax,
        "schedule": schedule,
        "duplicate_jobs": duplicate_jobs,
    }


def status_matches(source_status, file_marker):
    if source_status == "FINISHED":
        return file_marker == "LMAX_FIRST"
    if source_status in {"TIMEOUT", "TIME_LIMIT_FEASIBLE"}:
        return file_marker in {"TIMEOUT", "TIME_LIMIT_FEASIBLE", "LMAX_FIRST"}
    if source_status in {"UNSAT", "INFEASIBLE"}:
        return file_marker == source_status
    return file_marker == source_status


def validate_row(instance, display_solver, solver_id, source_row):
    n = instance["n"]
    instance_type = instance["type"]
    filename = instance["filename"]
    instance_path = INS_DIR / f"wtrd_pred{n}" / instance_type / filename
    solution_path = DATA_DIR / f"solutions_{solver_id}" / f"{n}-{instance_type}" / f"{filename}.txt"
    source_status = source_row["status"] if source_row else "MISSING_SOURCE_ROW"
    result_lmax = source_row["value"] if source_row else None
    parsed = parse_solution_file(solution_path)
    issues = []

    if source_row is None:
        issues.append("missing_source_row")
    if not instance_path.exists():
        issues.append("missing_instance_file")
    if not parsed["exists"]:
        issues.append("missing_solution_file")

    file_status_match = status_matches(source_status, parsed["marker"])
    if parsed["exists"] and not file_status_match:
        issues.append("file_status_mismatch")

    schedule = parsed["schedule"]
    schedule_valid = None
    violations = []
    recomputed_lmax = None
    if instance_path.exists() and schedule:
        n_read, durations, ready_dates, due_dates, deadlines, successors = read_dataset(instance_path)
        schedule_valid, violations = validate_schedule(
            schedule,
            n_read,
            durations,
            ready_dates,
            deadlines,
            successors,
        )
        if schedule_valid:
            recomputed_lmax = compute_max_lateness(
                schedule,
                durations,
                due_dates,
                clamp_zero=False,
            )
        else:
            issues.append("invalid_schedule")

    if parsed["duplicate_jobs"]:
        issues.append("duplicate_job_lines")

    if source_status == "FINISHED" and not schedule:
        issues.append("missing_finished_schedule")
    elif source_status in {"TIMEOUT", "TIME_LIMIT_FEASIBLE"} and not schedule:
        outcome = "NO_INCUMBENT"
    elif source_status in {"UNSAT", "INFEASIBLE"} and not schedule:
        outcome = "NO_SCHEDULE_EXPECTED"
    elif schedule and schedule_valid:
        outcome = "VALID_SCHEDULE"
    else:
        outcome = "INVALID"

    if recomputed_lmax is not None:
        if parsed["lmax"] is None:
            issues.append("missing_file_lmax")
        elif parsed["lmax"] != recomputed_lmax:
            issues.append("file_lmax_mismatch")
        if result_lmax is None:
            issues.append("missing_result_lmax")
        elif result_lmax != recomputed_lmax:
            issues.append("result_lmax_mismatch")

    return {
        "solver": display_solver,
        "solver_id": solver_id,
        "n": n,
        "type": instance_type,
        "filename": filename,
        "source_status": source_status,
        "file_marker": parsed["marker"],
        "status_match": file_status_match,
        "solution_exists": parsed["exists"],
        "schedule_jobs": len(schedule),
        "schedule_valid": schedule_valid,
        "recomputed_lmax": recomputed_lmax,
        "file_lmax": parsed["lmax"],
        "result_lmax": result_lmax,
        "outcome": outcome,
        "issues": "; ".join(issues),
        "violation_detail": "; ".join(violations[:5]),
        "solution_path": solution_path.relative_to(PROJECT_ROOT).as_posix(),
    }


def build_report():
    instances = read_filename_instances()
    source_rows = {
        display_solver: read_solver_results(workbook_path)
        for display_solver, (_, workbook_path) in SOLVER_SOURCES.items()
    }

    detail_rows = []
    for instance in instances:
        key = (instance["n"], instance["type"], instance["filename"])
        for display_solver, (solver_id, _) in SOLVER_SOURCES.items():
            detail_rows.append(
                validate_row(
                    instance,
                    display_solver,
                    solver_id,
                    source_rows[display_solver].get(key),
                )
            )

    details = pd.DataFrame(detail_rows)
    summary_rows = []
    for display_solver in SOLVER_SOURCES:
        solver_details = details[details["solver"] == display_solver]
        summary_rows.append(
            {
                "solver": display_solver,
                "solver_id": SOLVER_SOURCES[display_solver][0],
                "rows_checked": len(solver_details),
                "solution_files_found": int(solver_details["solution_exists"].sum()),
                "valid_schedules": int((solver_details["outcome"] == "VALID_SCHEDULE").sum()),
                "no_incumbent_timeouts": int((solver_details["outcome"] == "NO_INCUMBENT").sum()),
                "no_schedule_expected": int((solver_details["outcome"] == "NO_SCHEDULE_EXPECTED").sum()),
                "invalid_or_inconsistent": int((solver_details["issues"] != "").sum()),
                "status_mismatches": int((~solver_details["status_match"]).sum()),
                "lmax_mismatches": int(
                    solver_details["issues"].str.contains("lmax_mismatch", regex=False).sum()
                ),
            }
        )

    summary = pd.DataFrame(summary_rows)
    issues = details[details["issues"] != ""].copy()
    metadata = pd.DataFrame(
        [
            {"field": "instance_source", "value": "Filenames/"},
            {"field": "instance_count", "value": len(instances)},
            {"field": "solver_count", "value": len(SOLVER_SOURCES)},
            {"field": "rows_checked", "value": len(details)},
            {"field": "validation_rule", "value": "Raw ready/deadline windows, no overlap, precedence, true Lmax"},
            {"field": "timeout_rule", "value": "TIMEOUT with a schedule is validated as an incumbent, not as proven optimum"},
            {"field": "source_policy", "value": "No result workbook, solution file, or filename list is modified"},
        ]
    )
    return metadata, summary, details, issues


def style_report(path):
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    issue_fill = PatternFill("solid", fgColor="FCE4D6")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in worksheet.columns:
            letter = column[0].column_letter
            worksheet.column_dimensions[letter].width = min(
                60,
                max(12, max(len(str(cell.value or "")) for cell in column) + 2),
            )
        if worksheet.title == "Issues":
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.fill = issue_fill
    workbook.save(path)


def main():
    metadata, summary, details, issues = build_report()
    with pd.ExcelWriter(REPORT_WORKBOOK, engine="openpyxl") as writer:
        metadata.to_excel(writer, sheet_name="Metadata", index=False)
        summary.to_excel(writer, sheet_name="Summary", index=False)
        details.to_excel(writer, sheet_name="Details", index=False)
        issues.to_excel(writer, sheet_name="Issues", index=False)
    style_report(REPORT_WORKBOOK)
    print(f"Created validation report: {REPORT_WORKBOOK}")
    print(f"Checked {len(details)} solver-instance rows; issues={len(issues)}")


if __name__ == "__main__":
    main()
